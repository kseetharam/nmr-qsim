"""
fluo_suite_operators.py

Cluster-runnable generation of ZULF Lindblad jump operators + isotropic
Hamiltonian + NMR-protocol operators, for the candidate systems identified
in circ_sim/scripts/zulf_numerics/fluo_suite/notes/nuclei_selection_survey.tex
(Table 2, "Filtered cutoff selection").

Output format matches circ_sim/scripts/linblad_dyn/data/gemcitabine_operators.pkl
exactly (same metadata keys, same Pauli-string convention, same L_(k,m) /
L1_(k,m) naming), but built as Pauli-string dictionaries via
circ_sim/scripts/linblad_dyn/utils/pauli_algebra.py instead of hardcoded
per-molecule values, so it scales to the ~50-atom subsets in the table
(gemcitabine's own reference file only ever needed 10 spins).

Physics: L^{(1)}_{k,m} = sqrt(2 tau_c/3) Q^{(Z),1}_{k,m}    (9 ops, CSA only)
         L^{(2)}_{k,m} = sqrt(2 tau_c/5) (Q^{(Z),2}_{k,m} + Q_dip_{k,m})  (25 ops)
same formulas as circ_sim/scripts/linblad_dyn/utils/linblad_utils.py
(build_QZ_ops, build_Qdip_ops, build_jump_operators), ported to the
symbolic Pauli-string representation.

Atom selection reuses (imports directly, not re-derives) the chemistry
constraints and combined J+dipolar graph reachability from
circ_sim/scripts/zulf_numerics/fluo_suite/moment_convergence.py, so the
systems generated here are IDENTICAL to the ones tabulated in the survey.

KNOWN LIMITATION, recorded in every output file's metadata under the key
'j_coupling_coverage': the merged workbook only reports 1-bond/2-bond
J-couplings. Every other pair in a retained subset defaults to J=0 in H0 --
an approximation (unmeasured, not verified negligible), not a computed/
measured completeness. See the nuclei_selection_survey.tex discussion.

Assumptions fixed by the user for this run:
    B_vec = (0, 0, 5e-7) T
    tau_c = 1e-10 s for every system (not the MW-scaled heuristic discussed
            earlier -- explicitly overridden for this generation pass)

Usage
-----
    python fluo_suite_operators.py --list                 # enumerate all systems
    python fluo_suite_operators.py --index 17              # run system #17 (SLURM array index)
    python fluo_suite_operators.py --mol mol18 --anchor 8 --cutoff 2.5
    python fluo_suite_operators.py --all                   # run everything, sequentially (local testing only)
"""

import os
import sys
import argparse
import pickle
from itertools import combinations

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, 'utils'))
from pauli_algebra import PauliAlgebra                                            # noqa: E402

# moment_convergence.py owns the atom-selection methodology (chemistry
# constraints + combined J/dipolar graph reachability) that produced
# nuclei_selection_survey.tex; imported directly so this script can never
# silently drift from that table.
FLUO_SUITE_DIR = os.path.normpath(os.path.join(HERE, '..', 'zulf_numerics', 'fluo_suite'))
sys.path.insert(0, FLUO_SUITE_DIR)
from moment_convergence import (                                                   # noqa: E402
    WORKBOOK, load_molecule, list_cf_anchors, apply_chemistry_constraints,
    dipolar_weights, greedy_order,
)

GAMMA = {'F': 251.81520e6, 'H': 267.52218e6, 'C': 67.28284e6}   # rad/(s.T), 19F/1H/13C
HBAR = 1.054571817e-34   # J.s
MU0 = 1.25663706212e-6   # T.m/A
ANGSTROM = 1e-10         # m

B_VEC_T = (0.0, 0.0, 5e-7)   # fixed for this generation pass
TAU_C_S = 1e-10             # fixed for this generation pass (all systems)
SECOND_CUTOFFS = [3.0, 2.5, 2.0]   # scanned in this order, matching Table 2
CONNECTIVITY_FLOOR = 9

DATA_DIR = os.path.join(HERE, 'data', 'fluo_suite')

_RT2 = np.sqrt(2.0)
_RT3 = np.sqrt(3.0)
_RT6 = np.sqrt(6.0)

_CG_TABLE = {
    (0, 1, -1): 1.0 / _RT3, (0, 0, 0): -1.0 / _RT3, (0, -1, 1): 1.0 / _RT3,
    (1, 1, 0): 1.0 / _RT2, (1, 0, 1): -1.0 / _RT2, (1, 1, -1): 1.0 / _RT2, (1, 0, 0): 0.0,
    (1, -1, 1): -1.0 / _RT2, (1, 0, -1): 1.0 / _RT2, (1, -1, 0): -1.0 / _RT2,
    (2, 1, 1): 1.0,
    (2, 1, 0): 1.0 / _RT2, (2, 0, 1): 1.0 / _RT2,
    (2, 1, -1): 1.0 / _RT6, (2, 0, 0): 2.0 / _RT6, (2, -1, 1): 1.0 / _RT6,
    (2, 0, -1): 1.0 / _RT2, (2, -1, 0): 1.0 / _RT2,
    (2, -1, -1): 1.0,
}


def _cg1x1(l, m1, m2):
    return _CG_TABLE.get((l, m1, m2), 0.0)


def _b_spherical(b_vec):
    bx, by, bz = b_vec
    return {0: bz + 0j, +1: -(bx + 1j * by) / _RT2, -1: (bx - 1j * by) / _RT2}


def _sigma_lm(sigma_t):
    """Irreducible spherical components (l=0,1,2) of a general (possibly
    non-symmetric) 3x3 tensor. Same formulas as linblad_utils.py's _sigma_lm
    -- l=1 comes from the antisymmetric part, so it is NOT assumed to
    vanish (only turns out to for tensors that happen to be symmetric)."""
    s = np.asarray(sigma_t, dtype=complex)
    s_iso = np.trace(s) / 3.0
    return {
        (0, 0): -np.sqrt(3) * s_iso,
        (1, 0): -1j / _RT2 * (s[0, 1] - s[1, 0]),
        (1, +1): -0.5 * ((s[2, 0] - s[0, 2]) + 1j * (s[2, 1] - s[1, 2])),
        (1, -1): -0.5 * ((s[2, 0] - s[0, 2]) - 1j * (s[2, 1] - s[1, 2])),
        (2, 0): np.sqrt(2.0 / 3.0) * (s[2, 2] - s_iso),
        (2, +1): -0.5 * ((s[0, 2] + s[2, 0]) + 1j * (s[1, 2] + s[2, 1])),
        (2, -1): +0.5 * ((s[0, 2] + s[2, 0]) - 1j * (s[1, 2] + s[2, 1])),
        (2, +2): 0.5 * ((s[0, 0] - s[1, 1]) + 1j * (s[0, 1] + s[1, 0])),
        (2, -2): 0.5 * ((s[0, 0] - s[1, 1]) - 1j * (s[0, 1] + s[1, 0])),
    }


def _s_sph(alg, pos, q):
    if q == 0:
        return alg.iz(pos)
    if q == +1:
        return alg.scale(alg.ip(pos), -1 / _RT2)
    if q == -1:
        return alg.scale(alg.im(pos), 1 / _RT2)


def _t_lk(alg, pos, l, k, b_sph):
    """T^{(l)}_k(S, B) = sum_{q1+q2=k} <1,q1;1,q2|l,k> S_sph[q1] B_sph[q2]."""
    result = {}
    for q1 in (-1, 0, 1):
        q2 = k - q1
        if q2 not in (-1, 0, 1):
            continue
        c = _cg1x1(l, q1, q2)
        if abs(c) < 1e-15:
            continue
        bq2 = b_sph[q2]
        if abs(bq2) < 1e-30:
            continue
        result = alg.add(result, alg.scale(_s_sph(alg, pos, q1), c * bq2))
    return result


def _t2_pair(alg, pi, pj, k):
    """Rank-2 two-spin IST T^{(2)}_k(i,j) (i != j)."""
    if k == 0:
        return alg.scale(alg.add(alg.scale(alg.product(alg.iz(pi), alg.iz(pj)), 2),
                                 alg.scale(alg.product(alg.ix(pi), alg.ix(pj)), -1),
                                 alg.scale(alg.product(alg.iy(pi), alg.iy(pj)), -1)),
                         1 / _RT6)
    if k == +1:
        return alg.scale(alg.add(alg.product(alg.ip(pi), alg.iz(pj)),
                                 alg.product(alg.iz(pi), alg.ip(pj))), -0.5)
    if k == -1:
        return alg.scale(alg.add(alg.product(alg.im(pi), alg.iz(pj)),
                                 alg.product(alg.iz(pi), alg.im(pj))), 0.5)
    if k == +2:
        return alg.scale(alg.product(alg.ip(pi), alg.ip(pj)), 0.5)
    if k == -2:
        return alg.scale(alg.product(alg.im(pi), alg.im(pj)), 0.5)


def _a2m_dict(r_vec):
    rhat = r_vec / np.linalg.norm(r_vec)
    a = 3 * np.outer(rhat, rhat) - np.eye(3)
    return {
         0: (2 * a[2, 2] - a[0, 0] - a[1, 1]) / _RT6,
        +1: -(a[0, 2] - 1j * a[1, 2]),
        -1:  (a[0, 2] + 1j * a[1, 2]),
        +2:  (a[0, 0] - a[1, 1] - 2j * a[0, 1]) / 2,
        -2:  (a[0, 0] - a[1, 1] + 2j * a[0, 1]) / 2,
    }


# ---------------------------------------------------------------------------
# Data loading beyond what moment_convergence.load_molecule provides:
# full 3x3 shielding tensors (needed for CSA) and molecular weight (not
# needed now that tau_c is fixed at 1e-10 s for every system, but kept for
# provenance/metadata).
# ---------------------------------------------------------------------------

def load_shielding_tensors(mol_id):
    """dict {atom_idx: (3,3) ndarray}, raw (generally non-symmetric)
    lab-frame GIAO tensor in ppm, for every H/C/F atom of mol_id."""
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb['Shielding tensors (full 3x3)']
    tensors = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] != mol_id:
            continue
        idx = row[2]
        sxx, sxy, sxz, syx, syy, syz, szx, szy, szz = row[4:13]
        tensors[idx] = np.array([[sxx, sxy, sxz], [syx, syy, syz], [szx, szy, szz]], dtype=float)
    return tensors


def load_molecular_weight(mol_id):
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb['Overview']
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] == mol_id:
            return row[4]
    return None


def load_drug_name(mol_id):
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb['Overview']
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] == mol_id:
            return row[1]
    return None


# ---------------------------------------------------------------------------
# System manifest: reproduces Table 2 ("Filtered cutoff selection") of
# nuclei_selection_survey.tex exactly -- single source of truth for which
# (molecule, anchor, cutoff) combinations get generated.
# ---------------------------------------------------------------------------

def list_systems():
    """Return the full ordered list of systems to generate, one dict per
    system, each with a stable 0-based 'index' for SLURM array jobs:

        {index, mol_id, name, anchor_c, anchor_f, n_anchors, version_idx,
         cutoff_ang, n_atoms, atom_indices}

    For every molecule/anchor version whose 5.0-A connected count is
    >= CONNECTIVITY_FLOOR, this includes the 5.0-A system, PLUS the
    "second cutoff" system from Table 2 when one exists (the first of
    SECOND_CUTOFFS at which the connected count changes from its 5.0-A
    value, provided that changed value is also >= CONNECTIVITY_FLOOR).
    """
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb['Overview']
    mols = [(row[0], row[1]) for row in ws.iter_rows(min_row=5, max_row=25, values_only=True)]

    systems = []
    for mol_id, name in mols:
        atoms, edges, bond_order, coords_ang = load_molecule(mol_id)
        anchors = list_cf_anchors(atoms, edges, bond_order)
        n_anchors = len(anchors)
        for version_idx, (anchor_c, strength, anchor_f) in enumerate(anchors):
            seed, candidate_pool, excluded = apply_chemistry_constraints(
                atoms, edges, bond_order, anchor_override=anchor_c)

            dip5 = dipolar_weights(atoms, coords_ang, cutoff_ang=5.0)
            order5, _, _ = greedy_order(atoms, edges, seed=seed, candidate_pool=candidate_pool,
                                        dip_weights=dip5)
            v5 = len(order5)
            if v5 < CONNECTIVITY_FLOOR:
                continue

            systems.append({
                'mol_id': mol_id, 'name': name, 'anchor_c': anchor_c, 'anchor_f': anchor_f,
                'n_anchors': n_anchors, 'version_idx': version_idx,
                'cutoff_ang': 5.0, 'n_atoms': v5, 'atom_indices': order5,
            })

            for cutoff in SECOND_CUTOFFS:
                dip_c = dipolar_weights(atoms, coords_ang, cutoff_ang=cutoff)
                order_c, _, _ = greedy_order(atoms, edges, seed=seed, candidate_pool=candidate_pool,
                                             dip_weights=dip_c)
                if len(order_c) != v5:
                    if len(order_c) >= CONNECTIVITY_FLOOR:
                        systems.append({
                            'mol_id': mol_id, 'name': name, 'anchor_c': anchor_c,
                            'anchor_f': anchor_f, 'n_anchors': n_anchors, 'version_idx': version_idx,
                            'cutoff_ang': cutoff, 'n_atoms': len(order_c), 'atom_indices': order_c,
                        })
                    break

    for i, s in enumerate(systems):
        s['index'] = i
    return systems


def system_label(sys_rec):
    letter = chr(ord('a') + sys_rec['version_idx']) if sys_rec['n_anchors'] > 1 else ''
    return f"{sys_rec['mol_id']}{letter}_{sys_rec['cutoff_ang']:.1f}A"


# ---------------------------------------------------------------------------
# Operator construction (Pauli-string dictionaries; no dense matrices)
# ---------------------------------------------------------------------------

def build_all_operators(atoms, edges, coords_ang, sigma_full, atom_indices,
                        tau_c=TAU_C_S, b_vec=B_VEC_T):
    """H0, Sz_weighted, Sy_op, coil, L1_ops (9), L2_ops (25) for the given
    atom subset, as Pauli-string dictionaries. Same physics/normalization
    as linblad_utils.py's build_jump_operators; see module docstring."""
    elements = dict(atoms)
    n = len(atom_indices)
    atom_pos = {idx: i for i, idx in enumerate(atom_indices)}
    alg = PauliAlgebra(n)
    subset = set(atom_indices)

    gammas = [GAMMA[elements[idx]] for idx in atom_indices]
    sigma_tilde = [np.eye(3) + 1e-6 * sigma_full[idx] for idx in atom_indices]
    sigma_iso = [np.trace(st) / 3.0 for st in sigma_tilde]
    b_sph = _b_spherical(b_vec)

    # --- H0: isotropic J-coupling + isotropic chemical shift ---
    h0 = {}
    j_pairs_populated = 0
    for (i, j), j_hz in edges.items():
        if i in subset and j in subset:
            j_pairs_populated += 1
            c = 2 * np.pi * j_hz
            pi, pj = atom_pos[i], atom_pos[j]
            h0 = alg.add(h0,
                         alg.scale(alg.product(alg.ix(pi), alg.ix(pj)), c),
                         alg.scale(alg.product(alg.iy(pi), alg.iy(pj)), c),
                         alg.scale(alg.product(alg.iz(pi), alg.iz(pj)), c))
    for i, idx in enumerate(atom_indices):
        c = -gammas[i] * sigma_iso[i]
        h0 = alg.add(h0, alg.scale(alg.iz(i), c * b_vec[2]))
        if abs(b_vec[0]) > 0 or abs(b_vec[1]) > 0:
            h0 = alg.add(h0, alg.scale(alg.ix(i), c * b_vec[0]),
                         alg.scale(alg.iy(i), c * b_vec[1]))

    # --- Q^{(Z),l}_{k,m} = sum_i gamma_i (-1)^{m+1} sigma_{l,m}(i) T^{(l)}_{-k}(S_i,B) ---
    slm_all = [_sigma_lm(st) for st in sigma_tilde]

    def build_qz(l):
        result = {}
        for k in range(-l, l + 1):
            for m in range(-l, l + 1):
                phase = (-1) ** (m + 1)
                op = {}
                for i in range(n):
                    slm = slm_all[i][(l, m)]
                    if abs(slm) < 1e-30:
                        continue
                    t_neg = _t_lk(alg, i, l, -k, b_sph)
                    if not t_neg:
                        continue
                    op = alg.add(op, alg.scale(t_neg, gammas[i] * phase * slm))
                result[(k, m)] = op
        return result

    qz1 = build_qz(1)
    qz2 = build_qz(2)

    # --- Q_dip_{k,m} = sum_{i<j} b_ij a_{2,m}^{ij} T^{(2)}_k(i,j) ---
    def b_dip(i, j):
        r = np.linalg.norm(coords_ang[j] - coords_ang[i]) * ANGSTROM
        return -(MU0 / (4 * np.pi)) * GAMMA[elements[i]] * GAMMA[elements[j]] * HBAR / r ** 3

    pairs = list(combinations(atom_indices, 2))
    b_vals = {(i, j): b_dip(i, j) for i, j in pairs}
    a_vals = {(i, j): _a2m_dict((coords_ang[j] - coords_ang[i]) * ANGSTROM) for i, j in pairs}

    q_dip = {}
    for k in range(-2, 3):
        for m in range(-2, 3):
            op = {}
            for i, j in pairs:
                c = b_vals[(i, j)] * a_vals[(i, j)][m]
                if abs(c) < 1e-30:
                    continue
                op = alg.add(op, alg.scale(_t2_pair(alg, atom_pos[i], atom_pos[j], k), c))
            q_dip[(k, m)] = op

    scale1 = np.sqrt(2 * tau_c / 3)
    scale2 = np.sqrt(2 * tau_c / 5)
    l1_ops = {km: alg.scale(op, scale1) for km, op in qz1.items()}
    l2_ops = {km: alg.scale(alg.add(qz2[km], q_dip[km]), scale2) for km in qz2}

    # --- NMR protocol operators (sudden-transfer + hard-pulse ZULF) ---
    sz_weighted, sy_op, coil = {}, {}, {}
    for i, idx in enumerate(atom_indices):
        w = GAMMA[elements[idx]] / GAMMA['H']
        sz_weighted = alg.add(sz_weighted, alg.scale(alg.iz(i), w))
        sy_op = alg.add(sy_op, alg.scale(alg.iy(i), w))
        coil = alg.add(coil, alg.scale(alg.ip(i), w))

    n_pairs_total = n * (n - 1) // 2
    j_coverage = {
        'populated_pairs': j_pairs_populated, 'total_pairs': n_pairs_total,
        'fraction': j_pairs_populated / n_pairs_total if n_pairs_total else None,
        'note': ("J_hz sourced ONLY from 1-bond/2-bond ORCA/PySCF data. Every other pair in "
                 "this subset is assumed exactly J=0 in H0 -- an approximation (unmeasured), "
                 "not a computed/verified negligibility. Q_dip and Q_CSA are complete "
                 "(coordinate- and shielding-tensor-derived, available for every pair/atom)."),
    }

    return h0, sz_weighted, sy_op, coil, l1_ops, l2_ops, j_coverage


def build_system(sys_rec):
    """Full output dict for one system record from list_systems(), in the
    same schema as gemcitabine_operators.pkl."""
    mol_id = sys_rec['mol_id']
    atoms, edges, bond_order, coords_ang = load_molecule(mol_id)
    sigma_full = load_shielding_tensors(mol_id)
    mw = load_molecular_weight(mol_id)
    elements = dict(atoms)
    atom_indices = sys_rec['atom_indices']
    n = len(atom_indices)

    h0, sz_weighted, sy_op, coil, l1_ops, l2_ops, j_coverage = build_all_operators(
        atoms, edges, coords_ang, sigma_full, atom_indices, tau_c=TAU_C_S, b_vec=B_VEC_T)

    spin_order = [f"1{elements[idx]}" if elements[idx] == 'H' else
                  (f"19{elements[idx]}" if elements[idx] == 'F' else f"13{elements[idx]}")
                  for idx in atom_indices]

    letter = chr(ord('a') + sys_rec['version_idx']) if sys_rec['n_anchors'] > 1 else ''
    system_name = (f"{sys_rec['name']} ({mol_id}{letter}) -- {n}-spin subset, "
                   f"anchor C{sys_rec['anchor_c']}, cutoff {sys_rec['cutoff_ang']:.1f} A")

    output = {
        'metadata': {
            'system': system_name,
            'source_molecule_id': mol_id,
            'source_molecule_name': sys_rec['name'],
            'source_molecular_weight_g_per_mol': mw,
            'anchor_carbon': sys_rec['anchor_c'],
            'anchor_bonded_F': sys_rec['anchor_f'],
            'n_fluorinated_anchors_in_molecule': sys_rec['n_anchors'],
            'cutoff_ang': sys_rec['cutoff_ang'],
            'atom_indices_source_workbook': list(atom_indices),
            'spin_order': spin_order,
            'B_vec_T': list(B_VEC_T),
            'tau_c_s': TAU_C_S,
            'pauli_convention': (
                f"{n}-char strings from {{I,X,Y,Z}}^{n}, ordered as spin_order; "
                f"M = sum_P c_P (P[0] otimes ... otimes P[{n - 1}]); "
                f"c_P = Tr(M P) / 2^{n}"
            ),
            'H0_units': 'rad/s',
            'L_units': 'sqrt(rad/s)',
            'L_description': (
                "Rank-1 and rank-2 Lindblad jump operators: "
                "L^{(1)}_{k,m} = sqrt(2 tau_c/3) Q^{(Z),1}_{k,m} (9 ops, k,m in {-1,0,1}, CSA "
                "antisymmetric-tensor channel only -- zero iff the shielding tensor at every "
                "atom happens to be symmetric); "
                "L^{(2)}_{k,m} = sqrt(2 tau_c/5) (Q^{(Z),2}_{k,m} + Q_dip_{k,m}) (25 ops, k,m in "
                "{-2,...,2}); extreme-narrowing limit (omega*tau_c << 1)."
            ),
            'nmr_protocol': (
                "Sudden-transfer + hard-pulse ZULF experiment. "
                "weights w_i = gamma_i / gamma_1H. "
                "Sz_weighted = sum_i w_i Sz_i (pre-pulse state rho_sud); "
                "Sy_op = sum_i w_i Sy_i (pi/2 Y-pulse generator); "
                "rho0 = exp(-i pi/2 Sy_op) rho_sud exp(i pi/2 Sy_op); "
                "coil = sum_i w_i S+_i (quadrature detection observable, dimensionless); "
                "FID(t) = Tr(coil . rho(t))"
            ),
            'j_coupling_coverage': j_coverage,
            'generated_by': 'circ_sim/scripts/linblad_dyn/fluo_suite_operators.py',
        },
        'H0': h0,
        'Sz_weighted': sz_weighted,
        'Sy_op': sy_op,
        'coil': coil,
    }
    for (k, m), op in l1_ops.items():
        output[f'L1_({k},{m})'] = op
    for (k, m), op in l2_ops.items():
        output[f'L_({k},{m})'] = op

    return output


def save_system(output, sys_rec, out_dir=DATA_DIR):
    os.makedirs(out_dir, exist_ok=True)
    label = system_label(sys_rec)
    path = os.path.join(out_dir, f'{label}_operators.pkl')
    with open(path, 'wb') as f:
        pickle.dump(output, f, protocol=pickle.HIGHEST_PROTOCOL)
    return path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _find_system(systems, mol_id, anchor_c, cutoff_ang):
    for s in systems:
        if s['mol_id'] == mol_id and s['anchor_c'] == anchor_c and abs(s['cutoff_ang'] - cutoff_ang) < 1e-6:
            return s
    raise ValueError(f"No system matches mol_id={mol_id}, anchor_c={anchor_c}, cutoff_ang={cutoff_ang}. "
                     f"Run --list to see valid combinations.")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--list', action='store_true', help='Print all systems with their index, then exit')
    p.add_argument('--index', type=int, help='Run the Nth system from list_systems() (0-based; SLURM $SLURM_ARRAY_TASK_ID)')
    p.add_argument('--mol', type=str, help='Molecule id, e.g. mol18 (use with --anchor and --cutoff)')
    p.add_argument('--anchor', type=int, help='Anchor carbon atom index (from --list output)')
    p.add_argument('--cutoff', type=float, help='Cutoff in Angstrom, e.g. 2.5')
    p.add_argument('--all', action='store_true', help='Run every system sequentially (local testing only)')
    p.add_argument('--out-dir', type=str, default=DATA_DIR, help=f'Output directory (default: {DATA_DIR})')
    args = p.parse_args()

    systems = list_systems()

    if args.list:
        for s in systems:
            print(f"{s['index']:3d}  {system_label(s):16s}  {s['name']:24s}  "
                  f"anchor=C{s['anchor_c']} (F{','.join(map(str, s['anchor_f']))})  "
                  f"n_atoms={s['n_atoms']}")
        print(f"\n{len(systems)} systems total (indices 0..{len(systems) - 1})")
        return

    if args.index is not None:
        targets = [systems[args.index]]
    elif args.mol is not None:
        if args.anchor is None or args.cutoff is None:
            p.error('--mol requires --anchor and --cutoff')
        targets = [_find_system(systems, args.mol, args.anchor, args.cutoff)]
    elif args.all:
        targets = systems
    else:
        p.error('Specify one of --list, --index, --mol/--anchor/--cutoff, or --all')
        return

    for sys_rec in targets:
        label = system_label(sys_rec)
        print(f"Building {label} ({sys_rec['name']}, {sys_rec['n_atoms']} spins, "
              f"anchor C{sys_rec['anchor_c']}, cutoff {sys_rec['cutoff_ang']:.1f} A) ...", flush=True)
        output = build_system(sys_rec)
        path = save_system(output, sys_rec, out_dir=args.out_dir)
        cov = output['metadata']['j_coupling_coverage']
        print(f"  Saved -> {path}  "
              f"(J-coverage: {cov['populated_pairs']}/{cov['total_pairs']} pairs, "
              f"{cov['fraction']:.1%})", flush=True)


if __name__ == '__main__':
    main()
