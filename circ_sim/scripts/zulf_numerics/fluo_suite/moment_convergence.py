"""
moment_convergence.py

Prototype for graph-based nucleus selection (see fluo_suite discussion):
greedily grow a spin subset outward from the 19F "reporter" nuclei of a
molecule, ranking candidates by cumulative |J| to the already-included set,
and track convergence of the ZULF FID via low-order spectral moments

    M_n = Tr[ coil * (-i)^n [H, [H, ... [H, rho0] ... ]] ]      (n commutators)

computed symbolically as Pauli-string dictionaries (same convention as
circ_sim/scripts/linblad_dyn/gemcitabine_jump_operators.py), so the cost
scales with the sparsity of the J-coupling graph rather than with 2^N —
tractable up to the ~80-spin molecules in the fluo_suite set.

For subset sizes <= EXACT_DIAG_MAX_SPINS an exact QuTiP Hamiltonian is also
built and its own moments computed by direct matrix commutators, as a
correctness cross-check on the symbolic (scalable) computation. Exact
diagonalization is never attempted beyond that size.

Data source: circ_sim/data/fluo_suite/ORCA_NMR_summary_merged.xlsx
    - 'Atomic coordinates' -> which atoms are NMR-active (H/C/F) for a
      given molecule (all carbons are treated as 13C-labelled, matching
      the workbook's own convention)
    - 'J-couplings (1J & 2J)' -> 'Recommended J iso (Hz)' column, defining
      the (sparse: 1-bond/2-bond only) coupling graph

Output: convergence plot + printed report under
    circ_sim/scripts/zulf_numerics/fluo_suite/data/
"""

import os
from collections import Counter
import numpy as np
import openpyxl
import qutip as qt
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

MOL_ID = 'mol11'   # Gemcitabine
EXACT_DIAG_MAX_SPINS = 10
N_MOMENTS = 4
TAU_C = 1e-10        # s, rotational correlation time (matches gemcitabine_jump_operators.py)
DIP_CUTOFF_ANG = 5.0  # Angstrom, matches the workbook's own 'F neighbours (<5A)' convention

GAMMA = {'F': 251.81520e6, 'H': 267.52218e6, 'C': 67.28284e6}   # rad/(s.T), 19F/1H/13C
HBAR = 1.054571817e-34   # J.s
MU0 = 1.25663706212e-6   # T.m/A
ANGSTROM = 1e-10         # m

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBOOK = os.path.normpath(os.path.join(
    HERE, '..', '..', '..', 'data', 'fluo_suite', 'ORCA_NMR_summary_merged.xlsx'
))
DATA_DIR = os.path.join(HERE, 'data')


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------

def load_molecule(mol_id):
    """Return (atoms, edges, bond_order, coords_ang) for a molecule in the
    merged fluo_suite workbook.

    atoms      : list of (idx, element) for every H/C/F atom
    edges      : dict {(i, j): J_hz} using the 'Recommended J iso (Hz)'
                 column, i < j, for every 1J/2J pair reported for this
                 molecule
    bond_order : dict {(i, j): n} the through-bond distance (1 or 2) for
                 the same pairs, needed to identify "directly bonded"
                 (1-bond) C-F and C-H pairs for the natural-abundance /
                 exchangeable-proton chemistry constraints.
    coords_ang : dict {idx: np.array([x, y, z])} in Angstrom, for every
                 H/C/F atom -- used to build the through-space dipolar
                 coupling graph (unlike the J graph, this is available
                 for EVERY pair, not just 1-bond/2-bond ones).
    """
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)

    coords_ws = wb['Atomic coordinates']
    atoms, coords_ang = [], {}
    for row in coords_ws.iter_rows(min_row=5, values_only=True):
        if row[0] != mol_id or row[3] not in ('H', 'C', 'F'):
            continue
        idx = row[2]
        atoms.append((idx, row[3]))
        coords_ang[idx] = np.array([row[4], row[5], row[6]], dtype=float)

    j_ws = wb['J-couplings (1J & 2J)']
    edges, bond_order = {}, {}
    for row in j_ws.iter_rows(min_row=5, values_only=True):
        if row[0] != mol_id:
            continue
        idx_a, idx_b, n_bond, j_rec = row[4], row[6], row[7], row[10]
        i, j = sorted((idx_a, idx_b))
        edges[(i, j)] = float(j_rec)
        bond_order[(i, j)] = int(n_bond)

    return atoms, edges, bond_order, coords_ang


def dipolar_weights(atoms, coords_ang, cutoff_ang=DIP_CUTOFF_ANG, tau_c=TAU_C):
    """Through-space dipolar 'rate' graph, in the SAME units (rad/s) as the
    coherent J-graph after converting J to angular frequency (2*pi*J).

    w_dip(i,j) = b_ij^2 * (2*tau_c/5), with
        b_ij = -(mu0/4pi) * gamma_i * gamma_j * hbar / r_ij^3   (rad/s)
    using the identical (2*tau_c/5) spectral-density prefactor already
    hardcoded into the rank-2 jump operators in gemcitabine_jump_operators.py
    / linblad_utils.py, so this is a genuine rate (not just a heuristic
    rescaling of 1/r^3) directly comparable to 2*pi*|J_ij|.

    Unlike the J graph (capped at 1-bond/2-bond by the source data), this
    is computed for EVERY pair within cutoff_ang -- through-space coupling
    doesn't care about the bonding path, so it can reconnect atoms that
    have no J-coupling route at all (e.g. once intermediate carbons are
    excluded by the natural-13C-abundance constraint). cutoff_ang matches
    the workbook's own 'F neighbours (<5A)' convention: 1/r^3 decay makes
    anything past ~5A negligible, and the cutoff keeps the graph sparse
    (both physically justified and needed to keep the symbolic Pauli-string
    moments tractable for the ~80-spin molecules).
    """
    elements = dict(atoms)
    idx_list = [idx for idx, _ in atoms]
    cutoff_m = cutoff_ang * ANGSTROM
    weights = {}
    for a_i, a_j in ((idx_list[p], idx_list[q])
                     for p in range(len(idx_list)) for q in range(p + 1, len(idx_list))):
        r_vec = (coords_ang[a_j] - coords_ang[a_i]) * ANGSTROM
        r = np.linalg.norm(r_vec)
        if r > cutoff_m:
            continue
        b_ij = -(MU0 / (4 * np.pi)) * GAMMA[elements[a_i]] * GAMMA[elements[a_j]] * HBAR / r ** 3
        i, j = sorted((a_i, a_j))
        weights[(i, j)] = (b_ij ** 2) * (2 * tau_c / 5)
    return weights


# ---------------------------------------------------------------------------
# Symbolic Pauli-string operator algebra (scales with graph sparsity, not 2^N)
# ---------------------------------------------------------------------------

_PM = {
    ('I', 'I'): ('I', 1 + 0j), ('I', 'X'): ('X', 1 + 0j), ('I', 'Y'): ('Y', 1 + 0j), ('I', 'Z'): ('Z', 1 + 0j),
    ('X', 'I'): ('X', 1 + 0j), ('X', 'X'): ('I', 1 + 0j), ('X', 'Y'): ('Z', 1j), ('X', 'Z'): ('Y', -1j),
    ('Y', 'I'): ('Y', 1 + 0j), ('Y', 'X'): ('Z', -1j), ('Y', 'Y'): ('I', 1 + 0j), ('Y', 'Z'): ('X', 1j),
    ('Z', 'I'): ('Z', 1 + 0j), ('Z', 'X'): ('Y', 1j), ('Z', 'Y'): ('X', -1j), ('Z', 'Z'): ('I', 1 + 0j),
}


class PauliAlgebra:
    """Pauli-string dictionary algebra over a fixed N-spin register.

    A dense-vector/matrix representation of size 2^N is never constructed;
    "excluded" spins (not in the current subset) simply carry 'I' at their
    string position, so subsets of any size share the same N-length keys.
    """

    def __init__(self, n_total):
        self.n = n_total
        self.i_n = 'I' * n_total

    def _single(self, idx, char, coeff):
        return {self.i_n[:idx] + char + self.i_n[idx + 1:]: coeff}

    def ix(self, idx): return self._single(idx, 'X', 0.5 + 0j)
    def iy(self, idx): return self._single(idx, 'Y', 0.5 + 0j)
    def iz(self, idx): return self._single(idx, 'Z', 0.5 + 0j)

    def ip(self, idx):
        return {**self._single(idx, 'X', 0.5 + 0j), **self._single(idx, 'Y', 0.5j)}

    def im(self, idx):
        return {**self._single(idx, 'X', 0.5 + 0j), **self._single(idx, 'Y', -0.5j)}

    @staticmethod
    def dagger(d):
        """Hermitian conjugate of a Pauli-string dict. Each Pauli STRING is
        itself Hermitian (tensor product of Hermitian single-site I/X/Y/Z),
        so daggering a linear combination only conjugates the coefficients."""
        return {k: v.conjugate() for k, v in d.items()}

    @staticmethod
    def add(*dicts):
        r = {}
        for d in dicts:
            for k, v in d.items():
                r[k] = r.get(k, 0j) + v
        return {k: v for k, v in r.items() if abs(v) > 1e-12}

    @staticmethod
    def scale(d, c):
        if abs(c) < 1e-30:
            return {}
        return {k: v * c for k, v in d.items()}

    def product(self, d1, d2):
        r = {}
        for s1, c1 in d1.items():
            for s2, c2 in d2.items():
                chars, ph = [], 1 + 0j
                for ch1, ch2 in zip(s1, s2):
                    ch, p = _PM[(ch1, ch2)]
                    chars.append(ch)
                    ph *= p
                key = ''.join(chars)
                r[key] = r.get(key, 0j) + c1 * c2 * ph
        return {k: v for k, v in r.items() if abs(v) > 1e-12}

    def commutator(self, d1, d2):
        return self.add(self.product(d1, d2), self.scale(self.product(d2, d1), -1))

    def trace_product(self, d1, d2):
        """Normalized (dimension-independent) trace <d1 . d2> = Tr[d1.d2]/2^N,
        equal to the coefficient of the all-identity string in the product.

        This is the SAME convention documented in gemcitabine_jump_operators.py
        ("c_P = Tr(M P) / 2^N") and is essential here: an un-normalized trace
        picks up a spurious factor of 2 for every included spin a given term
        happens NOT to touch (that spin is a trace "spectator" for that one
        term), so the raw trace inflates by ~2x every time an atom is added
        to the subset regardless of whether it is physically relevant. The
        normalized trace has no such artifact and is what should be compared
        step-to-step for a genuine convergence signal.
        """
        prod = self.product(d1, d2)
        return prod.get(self.i_n, 0j)


# ---------------------------------------------------------------------------
# Physical operators for a given spin subset
# ---------------------------------------------------------------------------

def build_operators(alg, atom_pos, subset_idx, detect_idx, elements, edges):
    """H_iso and rho0 (post-pulse) over subset_idx; coil over detect_idx only.

    rho0 follows the sudden-transfer + hard-pulse ZULF protocol used in
    gemcitabine_jump_operators.py: rho_sud = Sz_weighted is NOT itself a
    valid FID initial state, since Tr[coil . (nested commutators of an
    M-conserving H)_sud] vanishes identically for ALL orders (coil raises
    total M by 1, rho_sud conserves M, and a Heisenberg H_iso conserves M
    exactly, so the trace has no matching M-sector at any order). A hard
    pulse rho0 = exp(-i pi/2 Sy_op) rho_sud exp(+i pi/2 Sy_op), with
    Sy_op = sum_i w_i Iy_i, is required to rotate weight into the
    transverse plane first. Because Sy_op is a sum of commuting single-site
    terms, this collective rotation factorises into independent per-spin
    rotations by angle theta_i = (pi/2) w_i (a proton-calibrated hard pulse
    tips lower-gamma nuclei by proportionally less) - a closed form, no
    matrix exponential needed:  Iz_i -> cos(theta_i) Iz_i + sin(theta_i) Ix_i.

    coil (the detection operator) is built ONLY from detect_idx (the fixed
    19F reporters), NOT the growing subset_idx: a fluorine-detected ZULF
    experiment always measures the same F channel regardless of how many
    other spins are included in the dynamics, so the "signal" a moment
    measures is "how much do the currently-included spins perturb what F
    sees", not "how big is the whole system's total magnetization" (which
    trivially grows every time an atom is added and would never converge).

    atom_pos : dict {atom_idx (raw ORCA index): position in alg's N-spin
               register}. Atom indices are NOT contiguous (O/N atoms are
               excluded from the H/C/F register), so this remapping is
               required before indexing into the Pauli-string algebra.
    elements : dict {atom_idx: 'H'/'C'/'F'}
    edges    : dict {(i, j): J_hz} for the FULL molecule; only pairs with
               both endpoints in subset_idx contribute.
    """
    subset = set(subset_idx)
    H = {}
    for (i, j), j_hz in edges.items():
        if i in subset and j in subset:
            c = 2 * np.pi * j_hz
            pos_i, pos_j = atom_pos[i], atom_pos[j]
            H = alg.add(H,
                        alg.scale(alg.product(alg.ix(pos_i), alg.ix(pos_j)), c),
                        alg.scale(alg.product(alg.iy(pos_i), alg.iy(pos_j)), c),
                        alg.scale(alg.product(alg.iz(pos_i), alg.iz(pos_j)), c))

    rho0 = {}
    for idx in subset_idx:
        w = GAMMA[elements[idx]] / GAMMA['H']
        pos = atom_pos[idx]
        theta = (np.pi / 2) * w
        rho0 = alg.add(rho0,
                       alg.scale(alg.iz(pos), w * np.cos(theta)),
                       alg.scale(alg.ix(pos), w * np.sin(theta)))

    coil = {}
    for idx in detect_idx:
        w = GAMMA[elements[idx]] / GAMMA['H']
        coil = alg.add(coil, alg.scale(alg.ip(atom_pos[idx]), w))

    return H, rho0, coil


def symbolic_moments(alg, H, rho0, coil, n_moments):
    """M_k = <coil . (-i)^k [H,[H,...[H,rho0]...]]> for k = 0..n_moments
    (normalized trace <A> = Tr[A]/2^N; see PauliAlgebra.trace_product)."""
    moments = []
    c_k = rho0
    for k in range(n_moments + 1):
        if k > 0:
            c_k = alg.commutator(H, c_k)
        moments.append(((-1j) ** k) * alg.trace_product(coil, c_k))
    return np.array(moments)


# ---------------------------------------------------------------------------
# Dipolar dissipator (Tier 2): promotes the coherent-only moments above to
# full Liouvillian moments M_n = <coil . L^n[rho0]>, L[rho] = -i[H,rho] + D[rho]
# ---------------------------------------------------------------------------

def _a2m_dict(r_vec):
    """Rank-2 orientational factors {m: complex} for displacement r_vec.
    Same convention as gemcitabine_jump_operators.py's _a2m_dict."""
    rhat = r_vec / np.linalg.norm(r_vec)
    A = 3 * np.outer(rhat, rhat) - np.eye(3)
    return {
         0: (2 * A[2, 2] - A[0, 0] - A[1, 1]) / np.sqrt(6),
        +1: -(A[0, 2] - 1j * A[1, 2]),
        -1:  (A[0, 2] + 1j * A[1, 2]),
        +2:  (A[0, 0] - A[1, 1] - 2j * A[0, 1]) / 2,
        -2:  (A[0, 0] - A[1, 1] + 2j * A[0, 1]) / 2,
    }


def _T2_pair(alg, pi, pj, k):
    """Rank-2 two-spin IST T^(2)_k(i,j) as a Pauli-string dict (positions
    pi, pj already remapped via atom_pos). Same formulas as
    gemcitabine_jump_operators.py's _T2ij."""
    if k == 0:
        return alg.scale(alg.add(alg.scale(alg.product(alg.iz(pi), alg.iz(pj)), 2),
                                 alg.scale(alg.product(alg.ix(pi), alg.ix(pj)), -1),
                                 alg.scale(alg.product(alg.iy(pi), alg.iy(pj)), -1)),
                         1 / np.sqrt(6))
    if k == +1:
        return alg.scale(alg.add(alg.product(alg.ip(pi), alg.iz(pj)),
                                 alg.product(alg.iz(pi), alg.ip(pj))), -0.5)
    if k == -1:
        return alg.scale(alg.add(alg.product(alg.im(pi), alg.iz(pj)),
                                 alg.product(alg.iz(pi), alg.im(pj))), 0.5)
    if k == +2:
        return alg.scale(alg.product(alg.ip(pi), alg.ip(pj)), 0.5)
    if k == -2:
        return alg.scale(alg.product(alg.im(pi), alg.im(pj)), 0.5)


def build_dipolar_jump_ops(alg, atom_pos, subset_idx, coords_ang, elements,
                            tau_c=TAU_C, cutoff_ang=DIP_CUTOFF_ANG):
    """25 rank-2 dipolar Lindblad jump operators L_(k,m), restricted to
    pairs both in subset_idx and within cutoff_ang -- same physics and
    normalization (sqrt(2*tau_c/5)) as gemcitabine_jump_operators.py's
    Q_dip + scale2, but only the dipolar (not CSA) mechanism, matching
    this discussion's focus. Empty (all-subset-disconnected) operators are
    dropped for efficiency.
    """
    cutoff_m = cutoff_ang * ANGSTROM
    pairs = []
    for p in range(len(subset_idx)):
        for q in range(p + 1, len(subset_idx)):
            i, j = subset_idx[p], subset_idx[q]
            r_vec_ang = coords_ang[j] - coords_ang[i]
            r_m = r_vec_ang * ANGSTROM
            r = np.linalg.norm(r_m)
            if r > cutoff_m:
                continue
            b_ij = -(MU0 / (4 * np.pi)) * GAMMA[elements[i]] * GAMMA[elements[j]] * HBAR / r ** 3
            pairs.append((i, j, b_ij, _a2m_dict(r_m)))

    scale2 = np.sqrt(2 * tau_c / 5)
    jump_ops = []
    for k in range(-2, 3):
        for m in range(-2, 3):
            op = {}
            for i, j, b_ij, a2m in pairs:
                c = b_ij * a2m[m]
                if abs(c) < 1e-30:
                    continue
                op = alg.add(op, alg.scale(_T2_pair(alg, atom_pos[i], atom_pos[j], k), c))
            op = alg.scale(op, scale2)
            if op:
                jump_ops.append(op)
    return jump_ops


def apply_dissipator(alg, jump_ops, rho):
    """D[rho] = sum_alpha (L_a rho L_a^dag - 1/2 {L_a^dag L_a, rho})."""
    D = {}
    for L in jump_ops:
        Ld = alg.dagger(L)
        LdL = alg.product(Ld, L)
        term1 = alg.product(alg.product(L, rho), Ld)
        anticomm = alg.add(alg.product(LdL, rho), alg.product(rho, LdL))
        D = alg.add(D, term1, alg.scale(anticomm, -0.5))
    return D


def apply_liouvillian(alg, H, jump_ops, rho):
    """L[rho] = -i[H, rho] + D[rho] (D = 0 if jump_ops is empty, recovering
    the pure-Hamiltonian case exactly)."""
    return alg.add(alg.scale(alg.commutator(H, rho), -1j), apply_dissipator(alg, jump_ops, rho))


def liouvillian_moments(alg, H, jump_ops, rho0, coil, n_moments):
    """M_k = <coil . L^k[rho0]> for k = 0..n_moments, L = -i[H,.] + D[.].
    With jump_ops=[] this is IDENTICAL (not just equivalent up to
    convention) to symbolic_moments -- the -i is baked into each
    Liouvillian application here instead of factored out as (-i)^k at the
    end, which are the same n-fold-nested-commutator result either way."""
    moments = []
    c_k = rho0
    for k in range(n_moments + 1):
        if k > 0:
            c_k = apply_liouvillian(alg, H, jump_ops, c_k)
        moments.append(alg.trace_product(coil, c_k))
    return np.array(moments)


# ---------------------------------------------------------------------------
# Exact (QuTiP) cross-check, capped at EXACT_DIAG_MAX_SPINS
# ---------------------------------------------------------------------------

def exact_moments_and_spectrum(subset_idx, detect_idx, elements, edges, n_moments):
    n = len(subset_idx)
    local = {idx: k for k, idx in enumerate(subset_idx)}

    def embed(op_char, k):
        ops = [qt.qeye(2)] * n
        ops[k] = qt.jmat(0.5, op_char)
        return qt.tensor(ops)

    Ix = [embed('x', k) for k in range(n)]
    Iy = [embed('y', k) for k in range(n)]
    Iz = [embed('z', k) for k in range(n)]
    Ip = [Ix[k] + 1j * Iy[k] for k in range(n)]

    H = 0 * Ix[0]
    for (i, j), j_hz in edges.items():
        if i in local and j in local:
            c = 2 * np.pi * j_hz
            ki, kj = local[i], local[j]
            H += c * (Ix[ki] * Ix[kj] + Iy[ki] * Iy[kj] + Iz[ki] * Iz[kj])

    rho0 = 0 * Ix[0]
    for idx in subset_idx:
        w = GAMMA[elements[idx]] / GAMMA['H']
        theta = (np.pi / 2) * w
        k = local[idx]
        rho0 += w * (np.cos(theta) * Iz[k] + np.sin(theta) * Ix[k])
    coil = sum((GAMMA[elements[idx]] / GAMMA['H']) * Ip[local[idx]] for idx in detect_idx)

    H_mat, rho0_mat, coil_mat = H.full(), rho0.full(), coil.full()
    moments = []
    c_k = rho0_mat
    for k in range(n_moments + 1):
        if k > 0:
            c_k = H_mat @ c_k - c_k @ H_mat
        # normalized trace <A> = Tr[A]/2^n, matching PauliAlgebra.trace_product
        moments.append(((-1j) ** k) * np.trace(coil_mat @ c_k) / (2 ** n))

    evals = H.eigenenergies() / (2 * np.pi)   # Hz
    return np.array(moments), evals


def exact_liouvillian_moments(subset_idx, detect_idx, elements, edges, coords_ang, n_moments,
                               tau_c=TAU_C, cutoff_ang=DIP_CUTOFF_ANG):
    """Exact (QuTiP-matrix) cross-check for liouvillian_moments: same H,
    rho0, coil as exact_moments_and_spectrum, plus the 25 dipolar jump
    operator MATRICES (same physics as build_dipolar_jump_ops), and full
    Lindblad dissipator applied by direct matrix multiplication."""
    n = len(subset_idx)
    local = {idx: k for k, idx in enumerate(subset_idx)}

    def embed(op_char, k):
        ops = [qt.qeye(2)] * n
        ops[k] = qt.jmat(0.5, op_char)
        return qt.tensor(ops)

    Ix = [embed('x', k) for k in range(n)]
    Iy = [embed('y', k) for k in range(n)]
    Iz = [embed('z', k) for k in range(n)]
    Ip = [Ix[k] + 1j * Iy[k] for k in range(n)]
    Im = [Ix[k] - 1j * Iy[k] for k in range(n)]

    H = 0 * Ix[0]
    for (i, j), j_hz in edges.items():
        if i in local and j in local:
            c = 2 * np.pi * j_hz
            ki, kj = local[i], local[j]
            H += c * (Ix[ki] * Ix[kj] + Iy[ki] * Iy[kj] + Iz[ki] * Iz[kj])

    rho0 = 0 * Ix[0]
    for idx in subset_idx:
        w = GAMMA[elements[idx]] / GAMMA['H']
        theta = (np.pi / 2) * w
        k = local[idx]
        rho0 += w * (np.cos(theta) * Iz[k] + np.sin(theta) * Ix[k])
    coil = sum((GAMMA[elements[idx]] / GAMMA['H']) * Ip[local[idx]] for idx in detect_idx)

    def T2_pair(ki, kj, k):
        if k == 0:
            return (2 * Iz[ki] * Iz[kj] - Ix[ki] * Ix[kj] - Iy[ki] * Iy[kj]) / np.sqrt(6)
        if k == +1:
            return -(Ip[ki] * Iz[kj] + Iz[ki] * Ip[kj]) / 2
        if k == -1:
            return (Im[ki] * Iz[kj] + Iz[ki] * Im[kj]) / 2
        if k == +2:
            return Ip[ki] * Ip[kj] / 2
        if k == -2:
            return Im[ki] * Im[kj] / 2

    cutoff_m = cutoff_ang * ANGSTROM
    pairs = []
    for p in range(n):
        for q in range(p + 1, n):
            i, j = subset_idx[p], subset_idx[q]
            r_m = (coords_ang[j] - coords_ang[i]) * ANGSTROM
            r = np.linalg.norm(r_m)
            if r > cutoff_m:
                continue
            b_ij = -(MU0 / (4 * np.pi)) * GAMMA[elements[i]] * GAMMA[elements[j]] * HBAR / r ** 3
            pairs.append((local[i], local[j], b_ij, _a2m_dict(r_m)))

    scale2 = np.sqrt(2 * tau_c / 5)
    jump_mats = []
    for k in range(-2, 3):
        for m in range(-2, 3):
            op = 0 * Ix[0]
            any_term = False
            for ki, kj, b_ij, a2m in pairs:
                c = b_ij * a2m[m]
                if abs(c) < 1e-30:
                    continue
                op += c * T2_pair(ki, kj, k)
                any_term = True
            if any_term:
                jump_mats.append((scale2 * op).full())

    H_mat, rho0_mat, coil_mat = H.full(), rho0.full(), coil.full()

    def liouvillian(rho_mat):
        out = -1j * (H_mat @ rho_mat - rho_mat @ H_mat)
        for L in jump_mats:
            Ld = L.conj().T
            LdL = Ld @ L
            out += L @ rho_mat @ Ld - 0.5 * (LdL @ rho_mat + rho_mat @ LdL)
        return out

    moments = []
    c_k = rho0_mat
    for k in range(n_moments + 1):
        if k > 0:
            c_k = liouvillian(c_k)
        moments.append(np.trace(coil_mat @ c_k) / (2 ** n))
    return np.array(moments)


# ---------------------------------------------------------------------------
# Greedy graph-based growth from the 19F seed
# ---------------------------------------------------------------------------

def greedy_order(atoms, edges, seed=None, candidate_pool=None, dip_weights=None):
    """Return (order, disconnected, dipolar_only) where `order` lists atom
    indices grown outward from `seed` (default: the F atoms) by cumulative
    coupling strength to the current set, restricted to `candidate_pool`
    (default: every non-seed atom). `disconnected` lists candidate-pool
    atoms never reachable through a nonzero edge.

    Growth score combines the coherent J channel (converted to rad/s via
    2*pi*|J_ij|) and, if `dip_weights` is given, the through-space dipolar
    channel (already in rad/s -- see dipolar_weights()) in matched units.
    `dipolar_only` lists atoms whose growth step had ZERO coherent (J)
    contribution to the set at the time they were added -- i.e. atoms
    reconnected purely by the dissipative/dipolar channel, which the J
    graph alone (capped at 1-bond/2-bond data) could never reach.
    """
    elements = dict(atoms)
    all_idx = set(elements)
    included = list(seed) if seed is not None else [idx for idx in all_idx if elements[idx] == 'F']
    order = list(included)
    included_set = set(included)
    remaining = (set(candidate_pool) if candidate_pool is not None
                 else all_idx - included_set) - included_set
    dip_weights = dip_weights or {}

    def coupling_to_set(d, cand, cur_set):
        total = 0.0
        for j in cur_set:
            key = tuple(sorted((cand, j)))
            if key in d:
                total += abs(d[key])
        return total

    dipolar_only = []
    while remaining:
        scored = []
        for cand in remaining:
            j_part = 2 * np.pi * coupling_to_set(edges, cand, included_set)
            dip_part = coupling_to_set(dip_weights, cand, included_set)
            scored.append((j_part + dip_part, j_part, dip_part, cand))
        scored.sort(key=lambda t: (-t[0], t[3]))
        best_total, best_j, best_dip, best_cand = scored[0]
        if best_total == 0.0:
            break
        if best_j == 0.0 and best_dip > 0.0:
            dipolar_only.append(best_cand)
        order.append(best_cand)
        included_set.add(best_cand)
        remaining.discard(best_cand)

    return order, sorted(remaining), dipolar_only


def list_cf_anchors(atoms, edges, bond_order):
    """Return all candidate 13C anchor carbons: every carbon with at least
    one 1-bond C-F coupling, ranked by cumulative 1-bond |J(C,F)| (summed
    over every F it's directly bonded to), strongest first.

    A molecule with a single fluorinated group (e.g. one CF2 or CF3) has
    exactly one anchor; a molecule with several chemically separate
    fluorinated anchors (e.g. two independent CF2 groups) has one entry
    per anchor -- since natural-abundance 13C means only one carbon is
    isotopically active per molecule, each anchor represents a distinct,
    equally-valid "which carbon happens to be the 13C" scenario, not just
    a single best guess.

    Returns a list of (carbon_idx, cumulative_J, bonded_F_list), sorted by
    cumulative_J descending.
    """
    elements = dict(atoms)
    c_atoms = [idx for idx, el in atoms if el == 'C']

    def is_cf_1bond(i, j):
        return ({elements[i], elements[j]} == {'C', 'F'} and bond_order.get((i, j)) == 1)

    anchors = []
    for c in c_atoms:
        bonded_f = sorted(j if i == c else i for (i, j) in edges if is_cf_1bond(i, j) and c in (i, j))
        if bonded_f:
            total = sum(abs(edges[tuple(sorted((c, f)))]) for f in bonded_f)
            anchors.append((c, total, bonded_f))
    anchors.sort(key=lambda t: -t[1])
    return anchors


def apply_chemistry_constraints(atoms, edges, bond_order, anchor_override=None):
    """Restrict the growth seed/candidate pool per two standard heuristics
    used when simulating ZULF spectra of natural-abundance samples:

    1. Natural 13C abundance (~1.1%) means at most ONE carbon is
       isotopically active in any given molecule at a time, so only a
       single representative carbon is retained. By default this is the
       carbon with the strongest cumulative 1-bond C-F coupling (see
       list_cf_anchors); pass `anchor_override` (a carbon index from
       list_cf_anchors) to instead build the constrained set for a
       SPECIFIC anchor -- used to generate one "version" of a molecule
       per distinct fluorinated anchor when there is more than one. All
       OTHER carbons are excluded from the candidate pool entirely (not
       just capped at one by the greedy search); the seed still includes
       EVERY 19F in the molecule regardless of which anchor is chosen,
       matching an actual F-detected experiment (all F resonances are
       observed simultaneously; only the splitting pattern near the
       labelled anchor differs between versions).
    2. Exchangeable protons (bonded to O/N: -OH, -NH2, ...) undergo fast
       solvent exchange and do not show resolved coherent J-coupling in
       practice, so they are excluded from the candidate pool. Identified
       here as any H with NO 1-bond C-H edge (i.e. not directly bonded to
       a carbon) -- a data-driven proxy for "bonded to O or N instead".
       (Anchor-independent: unaffected by anchor_override.)

    Returns (seed, candidate_pool, excluded) where excluded records WHY
    each atom was dropped, for reporting.
    """
    elements = dict(atoms)
    f_atoms = [idx for idx, el in atoms if el == 'F']
    c_atoms = [idx for idx, el in atoms if el == 'C']
    h_atoms = [idx for idx, el in atoms if el == 'H']

    anchors = list_cf_anchors(atoms, edges, bond_order)
    if not anchors:
        raise ValueError("No carbon is 1-bond coupled to any F atom for this molecule")
    anchor_carbons = [c for c, _, _ in anchors]
    if anchor_override is not None:
        if anchor_override not in anchor_carbons:
            raise ValueError(f"anchor_override={anchor_override} is not a valid C-F anchor "
                              f"(candidates: {anchor_carbons})")
        bonded_c = anchor_override
    else:
        bonded_c = anchor_carbons[0]
    other_c = [c for c in c_atoms if c != bonded_c]

    def is_ch_1bond(i, j):
        return ({elements[i], elements[j]} == {'C', 'H'} and bond_order.get((i, j)) == 1)

    non_exchangeable_h = {idx for idx in h_atoms
                          if any(is_ch_1bond(i, j) and idx in (i, j) for (i, j) in edges)}
    exchangeable_h = [h for h in h_atoms if h not in non_exchangeable_h]

    seed = f_atoms + [bonded_c]
    candidate_pool = sorted(non_exchangeable_h)
    excluded = {
        'other carbons (natural 13C abundance: only 1 retained)': other_c,
        'exchangeable H (bonded to O/N, not C)': exchangeable_h,
    }
    return seed, candidate_pool, excluded


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    atoms, edges, bond_order, coords_ang = load_molecule(MOL_ID)
    elements = dict(atoms)
    n_total = len(atoms)
    alg = PauliAlgebra(n_total)
    atom_pos = {idx: pos for pos, (idx, _) in enumerate(atoms)}
    dip_weights = dipolar_weights(atoms, coords_ang)

    seed, candidate_pool, excluded = apply_chemistry_constraints(atoms, edges, bond_order)
    order, disconnected, dipolar_only = greedy_order(
        atoms, edges, seed=seed, candidate_pool=candidate_pool, dip_weights=dip_weights)
    detect_idx = [idx for idx, el in atoms if el == 'F']

    print(f"Molecule {MOL_ID}: {n_total} NMR-active nuclei "
          f"({sum(1 for e in elements.values() if e == 'F')} F, "
          f"{sum(1 for e in elements.values() if e == 'C')} C, "
          f"{sum(1 for e in elements.values() if e == 'H')} H), "
          f"{len(edges)} J-coupling edges (1J/2J only), "
          f"{len(dip_weights)} dipolar edges (<= {DIP_CUTOFF_ANG:.1f} A)")
    print(f"  Chemistry-constrained seed: {[f'{elements[a]}{a}' for a in seed]}")
    for reason, idx_list in excluded.items():
        if idx_list:
            print(f"  Excluded ({reason}): {[f'{elements[a]}{a}' for a in idx_list]}")
    print(f"  Candidate pool for growth ({len(candidate_pool)} atoms): "
          f"{[f'{elements[a]}{a}' for a in candidate_pool]}")
    if dipolar_only:
        print(f"  Reconnected via dipolar channel ONLY (zero J-path to the set "
              f"at the time added): {[f'{elements[a]}{a}' for a in dipolar_only]}")
    if disconnected:
        print(f"  Disconnected even with the dipolar channel included "
              f"(zero predicted contribution to the F-detected spectrum): {disconnected}")
    print()

    rows = []
    prev_vec = None
    for n in range(len(seed), len(order) + 1):
        subset = order[:n]
        H, rho0, coil = build_operators(alg, atom_pos, subset, detect_idx, elements, edges)
        m_sym = symbolic_moments(alg, H, rho0, coil, N_MOMENTS)
        vec = np.abs(m_sym[1:])   # drop M_0 (just Tr[coil.rho0], not a dynamical moment)
        rel_change = (np.linalg.norm(vec - prev_vec) / np.linalg.norm(vec)
                      if prev_vec is not None and np.linalg.norm(vec) > 0 else np.nan)
        prev_vec = vec

        cross_check = None
        if n <= EXACT_DIAG_MAX_SPINS:
            m_exact, evals = exact_moments_and_spectrum(subset, detect_idx, elements, edges, N_MOMENTS)
            cross_check = np.max(np.abs(m_sym - m_exact))

        rows.append({
            'n': n, 'added': order[n - 1], 'element': elements[order[n - 1]],
            'moments': m_sym, 'rel_change': rel_change, 'cross_check_max_diff': cross_check,
        })

        cc_str = f"  cross-check max|diff|={cross_check:.3e}" if cross_check is not None else ""
        rc_str = f"{rel_change:.4f}" if not np.isnan(rel_change) else "   -  "
        print(f"  n={n:2d}  +{elements[order[n-1]]}{order[n-1]:<3d}  "
              f"|M1..M4|={np.round(vec, 4)}  rel_change={rc_str}{cc_str}")

    # --- suggested stopping point: first n with rel_change below tolerance,
    #     sustained for 2 consecutive steps ---
    tol = 0.03
    stop_n = None
    for k in range(1, len(rows) - 1):
        if rows[k]['rel_change'] is not None and not np.isnan(rows[k]['rel_change']):
            if rows[k]['rel_change'] < tol and rows[k + 1]['rel_change'] < tol:
                stop_n = rows[k]['n']
                break
    print()
    if stop_n:
        chosen = order[:stop_n]
        print(f"Suggested stopping point: n={stop_n} spins "
              f"(rel. moment change < {tol:.0%} for 2 consecutive additions)")
        print(f"  Selected atoms: {dict(Counter(elements[a] for a in chosen))} "
              f"= {[f'{elements[a]}{a}' for a in chosen]}")
    else:
        print(f"No stopping point found within tolerance {tol:.0%} over the full connected set.")
        chosen = order

    print()
    print("Comparison against the traditional heuristic (retain ALL "
          "non-exchangeable H + the 1 F-bonded C, no further trimming):")
    print(f"  Traditional set (n={len(order)}): "
          f"{dict(Counter(elements[a] for a in order))} = "
          f"{[f'{elements[a]}{a}' for a in order]}")
    if stop_n and stop_n < len(order):
        dropped = order[stop_n:]
        print(f"  Moment-convergence would additionally drop {len(order) - stop_n} atom(s) "
              f"beyond n={stop_n}: {[f'{elements[a]}{a}' for a in dropped]}")
    else:
        print("  Moment-convergence does not support trimming further than the "
              "traditional heuristic here (no atom beyond it changes the signal by < tol "
              "for 2 consecutive steps).")

    # -------------------------------------------------------------------
    # Tier 2: same growth order, but track full Liouvillian moments
    # (coherent H + dipolar dissipator) instead of pure-Hamiltonian ones.
    # -------------------------------------------------------------------
    print()
    print("=" * 70)
    print("Tier 2: coherent-only vs coherent+dipolar-dissipator moments")
    print("=" * 70)

    rows_diss = []
    prev_vec_diss = None
    for n in range(len(seed), len(order) + 1):
        subset = order[:n]
        H, rho0, coil = build_operators(alg, atom_pos, subset, detect_idx, elements, edges)
        jump_ops = build_dipolar_jump_ops(alg, atom_pos, subset, coords_ang, elements)
        m_diss = liouvillian_moments(alg, H, jump_ops, rho0, coil, N_MOMENTS)
        vec_diss = np.abs(m_diss[1:])
        rel_change_diss = (np.linalg.norm(vec_diss - prev_vec_diss) / np.linalg.norm(vec_diss)
                           if prev_vec_diss is not None and np.linalg.norm(vec_diss) > 0 else np.nan)
        prev_vec_diss = vec_diss

        cross_check_diss = None
        if n <= EXACT_DIAG_MAX_SPINS:
            m_exact_diss = exact_liouvillian_moments(subset, detect_idx, elements, edges,
                                                     coords_ang, N_MOMENTS)
            cross_check_diss = np.max(np.abs(m_diss - m_exact_diss))

        rows_diss.append({
            'n': n, 'rel_change': rel_change_diss, 'cross_check_max_diff': cross_check_diss,
            'n_jump_ops': len(jump_ops),
        })

        cc_str = (f"  cross-check max|diff|={cross_check_diss:.3e}"
                  if cross_check_diss is not None else "")
        rc_str = f"{rel_change_diss:.4f}" if not np.isnan(rel_change_diss) else "   -  "
        print(f"  n={n:2d}  +{elements[order[n-1]]}{order[n-1]:<3d}  "
              f"({len(jump_ops):2d} active dipolar ops)  "
              f"|M1..M4|={np.round(vec_diss, 4)}  rel_change={rc_str}{cc_str}")

    stop_n_diss = None
    for k in range(1, len(rows_diss) - 1):
        rc = rows_diss[k]['rel_change']
        if rc is not None and not np.isnan(rc) and rc < tol and rows_diss[k + 1]['rel_change'] < tol:
            stop_n_diss = rows_diss[k]['n']
            break

    print()
    print("Comparison: coherent-only vs coherent+dissipative stopping point")
    print(f"  Coherent-only         : stop at n={stop_n} "
          f"({dict(Counter(elements[a] for a in order[:stop_n])) if stop_n else 'n/a'})")
    print(f"  Coherent + dipolar D  : stop at n={stop_n_diss} "
          f"({dict(Counter(elements[a] for a in order[:stop_n_diss])) if stop_n_diss else 'n/a'})")
    if stop_n_diss and stop_n and stop_n_diss > stop_n:
        newly_relevant = order[stop_n:stop_n_diss]
        print(f"  Atoms the dissipative channel says to KEEP that the coherent-only "
              f"criterion had already dropped: {[f'{elements[a]}{a}' for a in newly_relevant]}")
    elif stop_n_diss == stop_n:
        print("  No change: the dipolar dissipator does not shift the stopping point here.")

    # --- plot ---
    os.makedirs(DATA_DIR, exist_ok=True)
    ns = [r['n'] for r in rows]
    rel_changes = [r['rel_change'] for r in rows]
    cc = [r['cross_check_max_diff'] for r in rows]
    ns_diss = [r['n'] for r in rows_diss]
    rel_changes_diss = [r['rel_change'] for r in rows_diss]
    cc_diss = [r['cross_check_max_diff'] for r in rows_diss]

    fig, ax1 = plt.subplots(figsize=(9, 5))
    ax1.plot(ns, rel_changes, 'o-', color='C0', label='coherent-only')
    ax1.plot(ns_diss, rel_changes_diss, 's--', color='C1', label='coherent + dipolar dissipator')
    ax1.axhline(tol, color='gray', ls='--', lw=1, label=f'tolerance ({tol:.0%})')
    ax1.axvline(EXACT_DIAG_MAX_SPINS, color='crimson', ls=':', lw=1,
                label=f'exact-diag cutoff (n={EXACT_DIAG_MAX_SPINS})')
    if stop_n:
        ax1.axvline(stop_n, color='C0', ls='-', lw=1.2, alpha=0.6, label=f'coherent stop (n={stop_n})')
    if stop_n_diss:
        ax1.axvline(stop_n_diss, color='C1', ls='-', lw=1.2, alpha=0.6,
                    label=f'coherent+diss stop (n={stop_n_diss})')
    ax1.set_yscale('log')
    ax1.set_xlabel('number of spins in subset (F-seeded greedy growth)')
    ax1.set_ylabel('relative change in |M1..M4|')
    ax1.set_title(f'{MOL_ID} (Gemcitabine): coherent-only vs +dipolar-dissipator convergence')
    ax1.legend(fontsize=8)
    ax1.grid(alpha=0.3)

    fpath = os.path.join(DATA_DIR, f'{MOL_ID}_moment_convergence.png')
    plt.tight_layout()
    plt.savefig(fpath, dpi=150)
    plt.close()
    print(f"\nPlot saved -> {fpath}")

    cc_valid = [c for c in cc if c is not None]
    if cc_valid:
        print(f"Symbolic-vs-exact cross-check, coherent-only (n<={EXACT_DIAG_MAX_SPINS}): "
              f"max|diff| over all steps = {max(cc_valid):.3e}")
    cc_diss_valid = [c for c in cc_diss if c is not None]
    if cc_diss_valid:
        print(f"Symbolic-vs-exact cross-check, +dissipator (n<={EXACT_DIAG_MAX_SPINS}): "
              f"max|diff| over all steps = {max(cc_diss_valid):.3e}")


if __name__ == '__main__':
    main()
